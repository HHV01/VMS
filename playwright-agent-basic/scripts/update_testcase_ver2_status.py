import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb_path = r'c:\Users\VNTT\OneDrive - VNTT\Desktop\Backup\Dự án\VMS\Testcase chức năng hệ thống VMS version 2.xlsx'
wb = openpyxl.load_workbook(wb_path)
ws = wb['Smart Search ver2 ']

# Update row 91 (SS-RT-001) to Pass
ws.cell(row=91, column=12, value='Pass')
ws.cell(row=91, column=10, value='Hệ thống hiển thị cảnh báo đỏ dưới ô tìm kiếm: "Lỗi xác thực dữ liệu đầu vào. query: Chuỗi phải có tối đa 500 ký tự" khi bấm Tìm kiếm.')

wb.save(wb_path)
print('Updated SS-RT-001 to Pass in Testcase workbook!')
