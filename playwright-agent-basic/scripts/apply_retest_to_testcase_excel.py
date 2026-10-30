import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

# 1. Update Testcase version 2 workbook
wb_path = r'c:\Users\VNTT\OneDrive - VNTT\Desktop\Backup\Dự án\VMS\Testcase chức năng hệ thống VMS version 2.xlsx'
wb = openpyxl.load_workbook(wb_path)
ws = wb['Smart Search ver2 ']

# Row 52 (SS-FLT-009) -> Pass
ws.cell(row=52, column=12, value='Pass')
ws.cell(row=52, column=10, value='[Retest 22/10/2026 sau khi Dev fix] Khi chọn Từ ngày > Đến ngày, nút Tìm kiếm lập tức tự động bị Disable (disabled:pointer-events-none), không cho phép gửi request.')

# Row 61 (SS-FLT-018) -> Pass
ws.cell(row=61, column=12, value='Pass')
ws.cell(row=61, column=10, value='[Retest 22/10/2026 sau khi Dev fix] Lọc theo biển số đã áp dụng điều kiện chặt chẽ. Khi không có xe khớp biển số, hệ thống trả về đúng 0 kết quả và Empty State, không còn trả xe sai biển hay xe không biển.')

# Row 59 (SS-FLT-016) -> Fail
ws.cell(row=59, column=12, value='Fail')
ws.cell(row=59, column=10, value='[Retest 22/10/2026] VẪN CÒN LỖI: Bảng thuộc tính đối tượng Phương tiện chỉ có Loại xe, Biển số, Màu sắc. Hoàn toàn chưa có trường "Hướng đi" theo đặc tả PSD 10.3.5.')

wb.save(wb_path)
print('Successfully updated Testcase version 2 workbook!')

# Check counts
passes = sum(1 for r in range(12, ws.max_row+1) if ws.cell(row=r, column=12).value == 'Pass')
fails = sum(1 for r in range(12, ws.max_row+1) if ws.cell(row=r, column=12).value == 'Fail')
print(f'Testcase counts: Total={passes+fails}, Pass={passes} ({(passes/(passes+fails)*100):.1f}%), Fail={fails} ({(fails/(passes+fails)*100):.1f}%)')
