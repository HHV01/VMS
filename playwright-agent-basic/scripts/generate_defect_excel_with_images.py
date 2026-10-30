import os
import sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image
from PIL import Image as PILImage

sys.stdout.reconfigure(encoding='utf-8')

# 1 Remaining Confirmed Defect after Dev fixes
defects_data = [
    {
        'stt': 1,
        'defect_id': 'SS_BUG_01',
        'tc_id': 'SS-FLT-016',
        'category': 'Giao diện (Frontend)',
        'severity': 'Medium',
        'title': 'Bảng thuộc tính đối tượng Phương tiện thiếu trường "Hướng đi"',
        'pre': 'Đang mở panel Bộ lọc nâng cao.',
        'steps': '1. Mở nhóm "Đặc điểm đối tượng" trong Bộ lọc nâng cao.\n2. Bấm "+ Thêm đối tượng", chọn "Phương tiện".\n3. Quan sát danh mục các thuộc tính lọc của phương tiện hiển thị trên panel.',
        'actual': 'Giao diện chỉ hiển thị 3 trường: Loại phương tiện, Biển số, Màu sắc. Hoàn toàn KHÔNG CÓ trường lựa chọn "Hướng đi" (Trái, Phải, Thẳng, Ngược chiều...).',
        'expected': 'Đối tượng Phương tiện bắt buộc phải hỗ trợ lựa chọn trường thuộc tính Hướng đi để phục vụ lọc xe di chuyển theo hướng xác định theo đặc tả PSD 10.3.5.',
        'ref': 'PSD 10.3.5 / Phương tiện',
        'img_path': 'evidence_defects/SS_BUG_FLT_016_thieu_huong_di.png'
    }
]

print('Initializing Workbook...')
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Báo Cáo Defect Smart Search'

# Page setup & Gridlines
ws.views.sheetView[0].showGridLines = True

# Colors
navy_blue = '1B365D'
header_blue = '2F5597'
light_blue = 'D9E1F2'
white = 'FFFFFF'
high_red_fill = 'FCE4D6'
high_red_text = 'C00000'
med_orange_fill = 'FFF2CC'
med_orange_text = 'B25900'
low_green_fill = 'E2EFDA'
low_green_text = '375623'
gray_border = 'D9D9D9'

font_title = Font(name='Segoe UI', size=16, bold=True, color=white)
font_subtitle = Font(name='Segoe UI', size=10, italic=True, color=white)
font_header = Font(name='Segoe UI', size=10, bold=True, color=white)
font_cell = Font(name='Segoe UI', size=9, bold=False, color='000000')
font_bold = Font(name='Segoe UI', size=9, bold=True, color='000000')

fill_title = PatternFill(start_color=navy_blue, end_color=navy_blue, fill_type='solid')
fill_header = PatternFill(start_color=header_blue, end_color=header_blue, fill_type='solid')
fill_card = PatternFill(start_color='F2F4F8', end_color='F2F4F8', fill_type='solid')

thin_border = Border(
    left=Side(style='thin', color=gray_border),
    right=Side(style='thin', color=gray_border),
    top=Side(style='thin', color=gray_border),
    bottom=Side(style='thin', color=gray_border)
)
header_border = Border(
    left=Side(style='thin', color='FFFFFF'),
    right=Side(style='thin', color='FFFFFF'),
    top=Side(style='thin', color=header_blue),
    bottom=Side(style='medium', color=navy_blue)
)

# 1. Header Banner
ws.merge_cells('A1:L1')
cell_a1 = ws['A1']
cell_a1.value = 'BÁO CÁO TỔNG HỢP DEFECT & FLOW TÁI HIỆN - VMS SMART SEARCH LIVE'
cell_a1.font = font_title
cell_a1.fill = fill_title
cell_a1.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 40

ws.merge_cells('A2:L2')
cell_a2 = ws['A2']
cell_a2.value = 'Hệ thống: VMS Live (http://180.148.1.242:8004) | Build: 2.8.0 | Module: Smart Search | Ngày cập nhật: 22/10/2026 (Sau Retest Dev Fix) | Người lập: Senior QA Lead'
cell_a2.font = font_subtitle
cell_a2.fill = fill_title
cell_a2.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[2].height = 22

# 2. Summary KPI Cards (Row 4)
ws.row_dimensions[3].height = 10
ws.row_dimensions[4].height = 36

kpi_data = [
    ('B4:C4', 'TỔNG SỐ DEFECT', '1 Lỗi', '4472C4'),
    ('E4:F4', 'MỨC ĐỘ HIGH (NGHIÊM TRỌNG)', '0 Lỗi (0%)', high_red_text),
    ('H4:I4', 'MỨC ĐỘ MEDIUM (TRUNG BÌNH)', '1 Lỗi (100%)', med_orange_text),
    ('K4:L4', 'MỨC ĐỘ LOW (THẤP)', '0 Lỗi (0%)', low_green_text),
]

for rng, label, val, color_code in kpi_data:
    ws.merge_cells(rng)
    top_left = rng.split(':')[0]
    cell = ws[top_left]
    cell.value = f'{label}: {val}'
    cell.font = Font(name='Segoe UI', size=11, bold=True, color=color_code)
    cell.fill = fill_card
    cell.alignment = Alignment(horizontal='center', vertical='center')
    start_col = rng[0]
    end_col = rng.split(':')[1][0]
    for c in [start_col, end_col]:
        ws[f'{c}4'].border = thin_border

ws.row_dimensions[5].height = 12

# 3. Table Column Headers (Row 6)
headers = [
    ('A', 'STT', 6),
    ('B', 'Defect ID', 13),
    ('C', 'Test Case ID', 14),
    ('D', 'Phân loại', 20),
    ('E', 'Mức độ', 12),
    ('F', 'Tiêu đề Defect', 34),
    ('G', 'Tiền điều kiện', 26),
    ('H', 'Flow các bước tái hiện (Step-by-Step)', 45),
    ('I', 'Kết quả thực tế (Actual Result)', 38),
    ('J', 'Kết quả mong đợi (Expected Result)', 34),
    ('K', 'Tham chiếu đặc tả', 18),
    ('L', 'Hình ảnh bằng chứng (Evidence Screenshot)', 42),
]

ws.row_dimensions[6].height = 28
for col_letter, col_title, col_width in headers:
    cell = ws[f'{col_letter}6']
    cell.value = col_title
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = header_border
    ws.column_dimensions[col_letter].width = col_width

# 4. Fill Data Rows & Embed Images
current_row = 6
for d in defects_data:
    current_row += 1
    r = current_row
    ws.row_dimensions[r].height = 145

    ws[f'A{r}'].value = d['stt']
    ws[f'B{r}'].value = d['defect_id']
    ws[f'C{r}'].value = d['tc_id']
    ws[f'D{r}'].value = d['category']
    ws[f'E{r}'].value = d['severity']
    ws[f'F{r}'].value = d['title']
    ws[f'G{r}'].value = d['pre']
    ws[f'H{r}'].value = d['steps']
    ws[f'I{r}'].value = d['actual']
    ws[f'J{r}'].value = d['expected']
    ws[f'K{r}'].value = d['ref']

    # Alignment & Fonts
    for col_letter in ['A', 'B', 'C', 'D', 'E', 'K']:
        cell = ws[f'{col_letter}{r}']
        cell.font = font_bold if col_letter in ['B', 'C', 'E'] else font_cell
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    for col_letter in ['F', 'G', 'H', 'I', 'J']:
        cell = ws[f'{col_letter}{r}']
        cell.font = font_bold if col_letter == 'F' else font_cell
        cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        cell.border = thin_border

    # Severity Styling
    sev_cell = ws[f'E{r}']
    sev_cell.fill = PatternFill(start_color=med_orange_fill, end_color=med_orange_fill, fill_type='solid')
    sev_cell.font = Font(name='Segoe UI', size=9, bold=True, color=med_orange_text)

    # Column L (Evidence Image Cell)
    cell_l = ws[f'L{r}']
    cell_l.border = thin_border
    cell_l.alignment = Alignment(horizontal='center', vertical='center')

    # Embed Screenshot
    img_file = d['img_path']
    if os.path.exists(img_file):
        try:
            pil_img = PILImage.open(img_file)
            os.makedirs('test-data', exist_ok=True)
            thumb_path = f'test-data/thumb_{d["defect_id"]}.png'
            pil_img.thumbnail((280, 140))
            pil_img.save(thumb_path)

            excel_img = Image(thumb_path)
            excel_img.width = pil_img.width
            excel_img.height = pil_img.height

            ws.add_image(excel_img, f'L{r}')
            print(f'Embedded image for {d["defect_id"]} in cell L{r} ({pil_img.width}x{pil_img.height})')
        except Exception as e:
            cell_l.value = f'[Lỗi nạp ảnh: {e}]'
    else:
        cell_l.value = f'[Ảnh không tồn tại: {img_file}]'

# Save ONLY to the single requested path
dest1 = r'c:\Users\VNTT\OneDrive - VNTT\Desktop\Backup\Dự án\VMS\Bao_cao_Defect_Smart_Search_VMS_kem_hinh_anh.xlsx'
wb.save(dest1)

print(f'\nSuccessfully generated professional Excel report with images at single location:\n{dest1}')
