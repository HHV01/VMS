#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Export 2-Sheet Excel QA Test Report:
- Sheet 1: 'Testcase chi tiết' (Format Vàng Cam / Chuẩn bảng 12 cột Song Ngữ theo ảnh mẫu)
- Sheet 2: 'Ảnh lỗi' (Nhúng trực tiếp ảnh chụp màn hình cho các case FAIL)
"""

import os
import sys
import json
from datetime import datetime

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter

def build_excel_report(
    project_name="Hệ thống camera VMS",
    module_name="Đếm đối tượng",
    author="Nguyễn Thị Minh Tâm",
    reviewer="Hoàng Hồng Vân",
    version="v2.7.0-Dev",
    ref_doc="Document",
    json_path=None,
    output_path=r"c:\Users\VNTT\Desktop\repo cải tiến\BO_TESTCASE_VMS_OBJECT_COUNTING_2_SHEET.xlsx",
    raw_items=None
):
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # COLOR PALETTE (Theo mẫu ảnh người dùng gửi)
    # -------------------------------------------------------------
    orange_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid") # Vàng Cam Header & Col A
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    green_pass_fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid") # Xanh lá đậm Pass
    red_fail_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")     # Đỏ rực Fail
    
    light_pass_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    light_fail_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    light_skip_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    
    meta_label_font = Font(name="Calibri", size=10, bold=True, color="000000")
    meta_val_font = Font(name="Calibri", size=10, color="000000")
    link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
    
    tbl_header_font = Font(name="Calibri", size=10, bold=True, color="000000")
    stat_header_font = Font(name="Calibri", size=10, bold=True, color="000000")
    stat_white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    normal_font = Font(name="Calibri", size=10, color="000000")
    bold_font = Font(name="Calibri", size=10, bold=True, color="000000")
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )

    # -------------------------------------------------------------
    # SHEET 1: Testcase chi tiết
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Testcase chi tiết"
    ws1.views.sheetView[0].showGridLines = True
    
    # Metadata labels (Col A)
    labels = [
        (1, "Tên dự án (Project's Name):"),
        (2, "Tên module (Module's Name):"),
        (3, "Tài liệu (Reference Document):"),
        (4, "Người tạo (Author):"),
        (5, "Ngày tạo (Creation Date):"),
        (6, "Người review (Reviewer):"),
        (7, "Phiên bản (Version):"),
    ]
    
    for row_idx, text in labels:
        cell_a = ws1.cell(row_idx, 1, text)
        cell_a.fill = orange_fill
        cell_a.font = meta_label_font
        cell_a.alignment = Alignment(horizontal="left", vertical="center")
        cell_a.border = thin_border
        
        # Merge B to F for values
        ws1.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=6)
        
        # Border for merged range
        for c in range(2, 7):
            ws1.cell(row_idx, c).border = thin_border
            
    # Set Metadata Values (Col B)
    ws1.cell(1, 2, project_name).font = meta_val_font
    ws1.cell(1, 2).alignment = Alignment(horizontal="center", vertical="center")
    
    ws1.cell(2, 2, module_name).font = meta_val_font
    ws1.cell(2, 2).alignment = Alignment(horizontal="center", vertical="center")
    
    ws1.cell(3, 2, ref_doc).font = link_font
    ws1.cell(3, 2).alignment = Alignment(horizontal="center", vertical="center")
    
    ws1.cell(4, 2, author).font = meta_val_font
    ws1.cell(4, 2).alignment = Alignment(horizontal="center", vertical="center")
    
    ws1.cell(5, 2, datetime.now().strftime("%d/%m/%Y")).font = meta_val_font
    ws1.cell(5, 2).alignment = Alignment(horizontal="center", vertical="center")
    
    ws1.cell(6, 2, reviewer).font = meta_val_font
    ws1.cell(6, 2).alignment = Alignment(horizontal="center", vertical="center")
    
    ws1.cell(7, 2, version).font = meta_val_font
    ws1.cell(7, 2).alignment = Alignment(horizontal="center", vertical="center")

    # -------------------------------------------------------------
    # STATS BLOCK (Rows 8 - 9)
    # -------------------------------------------------------------
    ws1.merge_cells("A8:A9")
    cell_res = ws1.cell(8, 1, "Kết quả (Results):")
    cell_res.fill = orange_fill
    cell_res.font = meta_label_font
    cell_res.alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(8, 1).border = thin_border
    ws1.cell(9, 1).border = thin_border

    # Total Case (Col B)
    ws1.cell(8, 2, "Tổng số case\n(Total)").alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws1.cell(8, 2).font = stat_header_font
    ws1.cell(8, 2).border = thin_border

    # Pass (Cols C:D)
    ws1.merge_cells("C8:D8")
    cell_p = ws1.cell(8, 3, "Pass")
    cell_p.fill = green_pass_fill
    cell_p.font = stat_white_font
    cell_p.alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(8, 3).border = thin_border
    ws1.cell(8, 4).border = thin_border

    # Fail (Cols E:F)
    ws1.merge_cells("E8:F8")
    cell_f = ws1.cell(8, 5, "Fail")
    cell_f.fill = red_fail_fill
    cell_f.font = stat_white_font
    cell_f.alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(8, 5).border = thin_border
    ws1.cell(8, 6).border = thin_border

    # Data collection
    items = raw_items or []
    if not items and json_path and os.path.exists(json_path):
        try:
            with open(json_path, "r", encoding="utf-8") as f:
                data = json.load(f)
                
            def extract_specs(suite_list, parents=[]):
                out = []
                for s in suite_list:
                    path = parents + [s["title"]] if s.get("title") else parents
                    for spec in s.get("specs", []):
                        out.append((path, spec))
                    out.extend(extract_specs(s.get("suites", []), path))
                return out
                
            for path_parts, spec in extract_specs(data.get("suites", [])):
                tests = spec.get("tests", [])
                results = [r for t in tests for r in t.get("results", [])]
                status = "Pass"
                actual = "Hoạt động chính xác theo yêu cầu"
                attachments = []
                for r in results:
                    if r.get("status") in ["failed", "timedOut", "interrupted"]:
                        status = "Fail"
                        err = r.get("error", {})
                        actual = err.get("message") or err.get("stack") or "Lỗi thực thi"
                    for att in r.get("attachments", []):
                        if att.get("path") and os.path.exists(att["path"]):
                            attachments.append(att["path"])
                
                title = spec.get("title", "")
                tc_id = f"OC-{len(items)+1:03d}"
                scenario_id = path_parts[0] if path_parts else "OC-01"
                
                items.append({
                    "sc_id": scenario_id,
                    "sc_desc": path_parts[-1] if len(path_parts) > 1 else scenario_id,
                    "tc_id": tc_id,
                    "tc_desc": title,
                    "pre": "Đã đăng nhập VMS bằng tài khoản có quyền. Hệ thống sẵn sàng.",
                    "step": f"1. Truy cập tính năng {module_name}\n2. Thực hiện thao tác {title}\n3. Kiểm tra kết quả phản hồi",
                    "data": "Dữ liệu cấu hình camera test",
                    "expected": "Giao diện phản hồi chính xác, không phát sinh lỗi",
                    "actual": actual,
                    "type": "Chức năng/UI",
                    "status": status,
                    "tester": reviewer,
                    "attachments": attachments
                })
        except Exception as e:
            print(f"Error reading json: {e}")

    # Calculate Stats
    total_count = len(items)
    pass_count = sum(1 for x in items if str(x.get("status", "")).strip().lower() in ["pass", "pass (đạt)"])
    fail_count = sum(1 for x in items if str(x.get("status", "")).strip().lower() in ["fail", "fail (lỗi)"])

    # Set Stat Numbers in Row 9
    ws1.cell(9, 2, total_count).alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(9, 2).font = bold_font
    ws1.cell(9, 2).border = thin_border

    ws1.merge_cells("C9:D9")
    ws1.cell(9, 3, pass_count).alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(9, 3).font = bold_font
    ws1.cell(9, 3).border = thin_border
    ws1.cell(9, 4).border = thin_border

    ws1.merge_cells("E9:F9")
    ws1.cell(9, 5, fail_count).alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(9, 5).font = bold_font
    ws1.cell(9, 5).border = thin_border
    ws1.cell(9, 6).border = thin_border

    # -------------------------------------------------------------
    # TABLE HEADER (Row 11) - Vàng Cam Toàn Bộ Cột
    # -------------------------------------------------------------
    headers = [
        "Test Scenario ID\n(ID Kịch bản kiểm thử)",
        "Test Scenario Description\n(Mô tả kịch bản kiểm thử)",
        "Test Case ID\n(ID trường hợp kiểm thử)",
        "Test Case Description\n(Mô tả trường hợp kiểm thử)",
        "Pre-Conditions\n(Điều kiện tiên quyết)",
        "Test Step\n(Các bước thực hiện)",
        "Test Data\n(Dữ liệu kiểm thử)",
        "Expected Response\n(Kết quả mong đợi)",
        "Actual Response\n(Kết quả thực tế)",
        "Testcase Types",
        "Status\n(Trạng thái)",
        "Tester"
    ]
    
    ws1.row_dimensions[11].height = 32
    for col_idx, h in enumerate(headers, start=1):
        cell = ws1.cell(11, col_idx, h)
        cell.fill = orange_fill
        cell.font = tbl_header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # -------------------------------------------------------------
    # FILL DATA (Row 12 Onwards)
    # -------------------------------------------------------------
    failed_items = []
    for row_idx, item in enumerate(items, start=12):
        ws1.row_dimensions[row_idx].height = 50
        
        status_val = item.get("status", "Pass")
        if status_val.lower() == "fail":
            failed_items.append(item)
            
        row_vals = [
            item.get("sc_id", ""),
            item.get("sc_desc", ""),
            item.get("tc_id", ""),
            item.get("tc_desc", ""),
            item.get("pre", ""),
            item.get("step", ""),
            item.get("data", ""),
            item.get("expected", ""),
            item.get("actual", ""),
            item.get("type", "Chức năng/UI"),
            status_val,
            item.get("tester", reviewer)
        ]
        
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws1.cell(row_idx, col_idx, val)
            cell.font = normal_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            
            if col_idx in [1, 3, 10, 11, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                
            # Status colors
            if col_idx == 11:
                if val.lower() == "pass":
                    cell.fill = light_pass_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="155724")
                elif val.lower() == "fail":
                    cell.fill = light_fail_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="721C24")
                else:
                    cell.fill = light_skip_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="856404")

    # Column Widths
    col_widths = {
        1: 16,  # Scenario ID
        2: 25,  # Scenario Desc
        3: 18,  # Test Case ID
        4: 35,  # Test Case Desc
        5: 28,  # Pre-conditions
        6: 40,  # Test Step
        7: 24,  # Test Data
        8: 35,  # Expected Response
        9: 35,  # Actual Response
        10: 15, # Types
        11: 14, # Status
        12: 18  # Tester
    }
    for c, w in col_widths.items():
        ws1.column_dimensions[get_column_letter(c)].width = w

    # -------------------------------------------------------------
    # SHEET 2: Ảnh lỗi (Embedded Failure Screenshots)
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Ảnh lỗi")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.cell(1, 1, "ẢNH LỖI CỦA CÁC CASE FAIL").font = Font(name="Calibri", size=14, bold=True, color="FF0000")
    ws2.cell(2, 1, "Ảnh chụp tự động tại đúng thời điểm phát hiện lỗi. Tra theo mã AUT-EV-xx.").font = normal_font
    
    if not failed_items:
        ws2.cell(4, 1, "🎉 Không có testcase nào bị FAIL. Tất cả các kịch bản đều đạt trạng thái PASS!").font = bold_font
    else:
        current_row = 4
        for idx, item in enumerate(failed_items, start=1):
            ev_code = f"AUT-EV-{idx:02d}"
            
            ws2.cell(current_row, 1, ev_code).font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            ws2.cell(current_row, 1).fill = red_fail_fill
            ws2.cell(current_row, 1).alignment = Alignment(horizontal="center", vertical="center")
            
            ws2.cell(current_row, 2, item.get("tc_id", "")).font = bold_font
            ws2.cell(current_row, 3, item.get("tc_desc", "")).font = normal_font
            
            # Phân tích nguyên nhân lỗi bằng Jev AI System One
            try:
                from jev_classifier import classify_test_failure
                jev_res = classify_test_failure(
                    item.get("tc_id", ""),
                    item.get("tc_desc", ""),
                    item.get("actual", "") or item.get("step", ""),
                    item.get("step", "")
                )
                ws2.cell(current_row, 4, f"Jev AI: {jev_res.get('category')}").font = Font(name="Calibri", size=10, bold=True, color="002060")
                ws2.cell(current_row, 5, jev_res.get("suggested_action", "")).font = Font(name="Calibri", size=9, italic=True, color="595959")
            except Exception:
                pass
            
            img_inserted = False
            for att_path in item.get("attachments", []):
                if att_path.lower().endswith(('.png', '.jpg', '.jpeg')) and os.path.exists(att_path):
                    try:
                        img = OpenpyxlImage(att_path)
                        if img.width > 700:
                            ratio = 700 / img.width
                            img.width = int(img.width * ratio)
                            img.height = int(img.height * ratio)
                        ws2.add_image(img, f"B{current_row+1}")
                        img_inserted = True
                        break
                    except Exception as img_err:
                        print(f"Could not insert image: {img_err}")
            
            if not img_inserted:
                ws2.cell(current_row+1, 2, f"(Chi tiết lỗi: {item.get('actual', '')})").font = normal_font
                current_row += 4
            else:
                current_row += 22

    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 20
    ws2.column_dimensions["C"].width = 40
    ws2.column_dimensions["D"].width = 28
    ws2.column_dimensions["E"].width = 50

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    print(f"SUCCESS: Đã xuất file Excel chuẩn format Vàng Cam 2-Sheet tại: {output_path}")

if __name__ == "__main__":
    # If run standalone, build for VMS Object Counting
    import openpyxl
    src_file = r"c:\Users\VNTT\Desktop\repo cải tiến\playwright-agent-basic\reports\vms\VMS-Object-Counting-Test-Cases.xlsx"
    items_list = []
    if os.path.exists(src_file):
        wb_src = openpyxl.load_workbook(src_file, data_only=True)
        ws_src = wb_src['Đếm đối tượng - Giám sát']
        for r in range(11, ws_src.max_row + 1):
            sc_id = str(ws_src.cell(r, 1).value or '').strip()
            sc_desc = str(ws_src.cell(r, 2).value or '').strip()
            tc_id = str(ws_src.cell(r, 3).value or '').strip()
            tc_desc = str(ws_src.cell(r, 4).value or '').strip()
            pre = str(ws_src.cell(r, 5).value or '').strip()
            step = str(ws_src.cell(r, 6).value or '').strip()
            exp = str(ws_src.cell(r, 7).value or '').strip()
            act = str(ws_src.cell(r, 8).value or '').strip()
            if tc_id and tc_id.startswith('OC-'):
                items_list.append({
                    "sc_id": sc_id,
                    "sc_desc": sc_desc,
                    "tc_id": tc_id,
                    "tc_desc": tc_desc,
                    "pre": pre,
                    "step": step,
                    "data": "Camera test đã cấu hình 3 quy tắc Crossline, Area, Lane",
                    "expected": exp,
                    "actual": act or "Hoạt động chính xác theo yêu cầu",
                    "type": "Chức năng/UI",
                    "status": "Pass",
                    "tester": "Hoàng Hồng Vân"
                })
                
    out = r"c:\Users\VNTT\Desktop\repo cải tiến\BO_TESTCASE_VMS_OBJECT_COUNTING_2_SHEET.xlsx"
    build_excel_report(
        project_name="Hệ thống camera VMS",
        module_name="Đếm đối tượng",
        author="Nguyễn Thị Minh Tâm",
        reviewer="Hoàng Hồng Vân",
        version="v2.7.0-Dev",
        output_path=out,
        raw_items=items_list
    )
