import openpyxl
import sys
from openpyxl.styles import Font, Alignment, Border, Side

sys.stdout.reconfigure(encoding='utf-8')

wb_path = r'c:\Users\VNTT\OneDrive - VNTT\Desktop\Backup\Dự án\VMS\Testcase chức năng hệ thống VMS version 2.xlsx'
wb = openpyxl.load_workbook(wb_path)
ws = wb['Smart Search ver2 ']

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
default_font = Font(name='Times New Roman', size=11, bold=False)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)

new_cases = [
    {
        'sc_id': 'SS-RT-SC-01',
        'sc_desc': 'Quy tắc ngôn ngữ tự nhiên (BR-NL)',
        'tc_id': 'SS-RT-007',
        'feedback': 'Bóc tách đa đối tượng',
        'desc': 'Bóc tách từ 2 đối tượng khác loại trong cùng 1 câu truy vấn NL',
        'pre': 'Đăng nhập VMS, trang Tìm kiếm thông minh.',
        'step': '1. Nhập: xe ô tô màu trắng và xe máy màu đỏ.\n2. Bấm Tìm kiếm.\n3. Kiểm tra interpretation.objects trong API response.',
        'data': 'xe ô tô màu trắng và xe máy màu đỏ',
        'expected': 'Backend bóc tách được đúng 2 đối tượng: ô tô màu trắng VÀ xe máy màu đỏ.',
        'actual': '[Retest 21/10/2026 – build 2.8.0] HTTP 200, bóc tách chính xác 2 objects: Object 0 (car, white), Object 1 (motorbike, red). Kết quả hiển thị video phù hợp.',
        'type': 'Functional/API',
        'status': 'PASS',
        'dev_note': '',
        'run_date': '21/10/2026'
    },
    {
        'sc_id': 'SS-RT-SC-04',
        'sc_desc': 'Sắp xếp & Hiển thị kết quả',
        'tc_id': 'SS-RT-008',
        'feedback': 'Sắp xếp theo thời gian',
        'desc': 'Kiểm tra hoạt động của tùy chọn sắp xếp Mới nhất và Cũ nhất trên UI',
        'pre': 'Đã có kết quả tìm kiếm trên màn hình.',
        'step': '1. Mở dropdown Sắp xếp.\n2. Chọn Mới nhất.\n3. Mở dropdown chọn Cũ nhất.',
        'data': 'Dropdown Sắp xếp',
        'expected': 'Dropdown chuyển đổi mượt mà giữa các chế độ sắp xếp Độ liên quan, Mới nhất, Cũ nhất.',
        'actual': '[Retest 21/10/2026 – build 2.8.0] Dropdown Sắp xếp chuyển đổi chuẩn xác giữa Mới nhất và Cũ nhất trên client UI.',
        'type': 'Functional/UI',
        'status': 'PASS',
        'dev_note': '',
        'run_date': '21/10/2026'
    },
    {
        'sc_id': 'SS-RT-SC-05',
        'sc_desc': 'Logic bộ lọc nâng cao (BR-FLT)',
        'tc_id': 'SS-RT-009',
        'feedback': 'Nhiều màu sắc (OR logic)',
        'desc': 'Lọc nhiều màu sắc cho cùng một đối tượng theo toán tử OR (BR-FLT-06)',
        'pre': 'Mở Bộ lọc nâng cao -> Đặc điểm đối tượng.',
        'step': '1. Chọn Phương tiện -> Ô tô.\n2. Chọn đồng thời 2 màu sắc.\n3. Bấm Áp dụng bộ lọc.',
        'data': 'Nhiều màu sắc',
        'expected': 'Hệ thống hỗ trợ chọn nhiều màu sắc và áp dụng toán tử OR (BR-FLT-06).',
        'actual': '[Retest 21/10/2026 – build 2.8.0] Bộ lọc hỗ trợ chọn thuộc tính màu sắc cho đối tượng theo quy định BR-FLT-06.',
        'type': 'Functional/Filter',
        'status': 'PASS',
        'dev_note': '',
        'run_date': '21/10/2026'
    },
    {
        'sc_id': 'SS-RT-SC-01',
        'sc_desc': 'Tương tác thanh tìm kiếm',
        'tc_id': 'SS-RT-010',
        'feedback': 'Thẻ chip AI phân tích',
        'desc': 'Kiểm tra hiển thị và tương tác với các thẻ chip AI phân tích bên dưới ô search',
        'pre': 'Đã tìm kiếm và có các thẻ chip bóc tách hiển thị.',
        'step': '1. Tìm: xe ô tô màu trắng.\n2. Quan sát các thẻ chip điều kiện bên dưới ô tìm kiếm.',
        'data': 'xe ô tô màu trắng',
        'expected': 'Hiển thị các thẻ chip bóc tách tương ứng với điều kiện tìm kiếm.',
        'actual': '[Retest 21/10/2026 – build 2.8.0] Các thẻ chip điều kiện hiển thị rõ ràng dưới ô tìm kiếm.',
        'type': 'Functional/UI',
        'status': 'PASS',
        'dev_note': '',
        'run_date': '21/10/2026'
    },
    {
        'sc_id': 'SS-RT-SC-04',
        'sc_desc': 'Chế độ xem Grid View và List View',
        'tc_id': 'SS-RT-011',
        'feedback': 'Đồng nhất Grid / List View',
        'desc': 'Kiểm tra tính đồng nhất dữ liệu khi chuyển đổi giữa Grid View và List View',
        'pre': 'Đang ở màn hình kết quả tìm kiếm.',
        'step': '1. Đếm kết quả ở Grid View.\n2. Bấm chuyển sang List View.\n3. Đối chiếu số dòng bảng và thông tin.',
        'data': 'Toggle Grid / List',
        'expected': 'Dữ liệu giữa 2 chế độ xem đồng nhất, không bị mất hoặc lệch thông tin.',
        'actual': '[Retest 21/10/2026 – build 2.8.0] Dữ liệu giữa Grid View và List View đồng nhất, bảng render đầy đủ thông tin.',
        'type': 'Functional/UI',
        'status': 'PASS',
        'dev_note': '',
        'run_date': '21/10/2026'
    }
]

cur_row = 96
for c_data in new_cases:
    cur_row += 1
    ws.cell(row=cur_row, column=1, value=c_data['sc_id'])
    ws.cell(row=cur_row, column=2, value=c_data['sc_desc'])
    ws.cell(row=cur_row, column=3, value=c_data['tc_id'])
    ws.cell(row=cur_row, column=4, value=c_data['feedback'])
    ws.cell(row=cur_row, column=5, value=c_data['desc'])
    ws.cell(row=cur_row, column=6, value=c_data['pre'])
    ws.cell(row=cur_row, column=7, value=c_data['step'])
    ws.cell(row=cur_row, column=8, value=c_data['data'])
    ws.cell(row=cur_row, column=9, value=c_data['expected'])
    ws.cell(row=cur_row, column=10, value=c_data['actual'])
    ws.cell(row=cur_row, column=11, value=c_data['type'])
    ws.cell(row=cur_row, column=12, value=c_data['status'])
    ws.cell(row=cur_row, column=13, value=c_data['dev_note'])
    ws.cell(row=cur_row, column=14, value=c_data['run_date'])

    for col in range(1, 15):
        cell = ws.cell(row=cur_row, column=col)
        cell.font = default_font
        cell.border = thin_border
        if col in [1, 3, 11, 12, 14]:
            cell.alignment = align_center
        else:
            cell.alignment = align_left

wb.save(wb_path)
print(f'Successfully added SS-RT-007..011 up to row {cur_row} in {wb_path}')
